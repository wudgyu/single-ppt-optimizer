from pathlib import Path

import fitz
import pytest
from docx import Document

from app.services.parser_errors import ParserError
from app.services.parsers.docx_parser import parse_docx
from app.services.parsers.pdf_parser import parse_pdf, render_pdf_pages


def _create_docx(path: Path) -> None:
    document = Document()
    document.add_heading("运营进展", level=1)
    document.add_paragraph("温州数安港归集635亿条")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "数据规模"
    table.cell(1, 1).text = "635亿条"
    document.save(str(path))


def test_parse_docx_extracts_headings_paragraphs_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "supplement.docx"
    _create_docx(path)

    document = parse_docx(
        path,
        document_id="doc_02",
        filename=path.name,
        extracted_text_path="jobs/job_1/parsed/doc_02.json",
    )

    headings = [block for block in document.parsed_blocks if block.block_type == "heading"]
    paragraphs = [block for block in document.parsed_blocks if block.block_type == "paragraph"]
    tables = [block for block in document.parsed_blocks if block.block_type == "table"]
    assert headings[0].heading_level == 1
    assert paragraphs[0].text == "温州数安港归集635亿条"
    assert tables[0].table is not None
    assert tables[0].table.row_count == 2
    assert tables[0].source_anchor.table_id == "table_1"


def test_parse_docx_rejects_corrupted_file(tmp_path: Path) -> None:
    path = tmp_path / "supplement.docx"
    path.write_bytes(b"PK\x03\x04 not an office archive")

    with pytest.raises(ParserError) as exc_info:
        parse_docx(
            path,
            document_id="doc_02",
            filename=path.name,
            extracted_text_path="jobs/job_1/parsed/doc_02.json",
        )

    assert exc_info.value.code == "CORRUPTED_FILE"


def _create_searchable_pdf(path: Path) -> None:
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "635 billion records")
        document.save(str(path))


def _create_scanned_pdf(path: Path) -> None:
    with fitz.open() as document:
        document.new_page()
        document.save(str(path))


def test_parse_pdf_extracts_searchable_text(tmp_path: Path) -> None:
    path = tmp_path / "supplement.pdf"
    _create_searchable_pdf(path)

    document = parse_pdf(
        path,
        document_id="doc_03",
        filename=path.name,
        extracted_text_path="jobs/job_1/parsed/doc_03.json",
    )

    assert document.page_count == 1
    assert document.image_only_pages == []
    assert any("635 billion records" in block.text for block in document.parsed_blocks)
    assert document.parsed_blocks[0].source_anchor.page == 1


def test_parse_pdf_marks_scanned_pages_without_ocr(tmp_path: Path) -> None:
    path = tmp_path / "scanned.pdf"
    _create_scanned_pdf(path)

    document = parse_pdf(
        path,
        document_id="doc_04",
        filename=path.name,
        extracted_text_path="jobs/job_1/parsed/doc_04.json",
    )

    assert document.image_only_pages == [1]
    assert document.parsed_blocks == []


def test_render_pdf_pages_creates_previews(tmp_path: Path) -> None:
    path = tmp_path / "supplement.pdf"
    _create_searchable_pdf(path)

    rendered = render_pdf_pages(path, tmp_path / "previews")

    assert len(rendered) == 1
    assert rendered[0].is_file()
    assert rendered[0].stat().st_size > 0
