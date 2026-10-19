from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app.schemas.source import ShapeType
from app.services.parser_errors import ParserError
from app.services.parsers.pptx_parser import parse_pptx


def _create_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    text_box.text_frame.text = "温州数安港归集635亿条"
    text_box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    presentation.slides.add_slide(presentation.slide_layouts[1])
    presentation.save(str(path))


def test_parse_pptx_extracts_shapes_and_geometry(tmp_path: Path) -> None:
    path = tmp_path / "source.pptx"
    _create_pptx(path)

    document = parse_pptx(
        path,
        document_id="doc_01",
        filename=path.name,
        extracted_text_path="jobs/job_1/parsed/doc_01.json",
    )

    assert document.page_count == 2
    shape = next(
        shape for shape in document.parsed_slides[0].shapes if shape.text == "温州数安港归集635亿条"
    )
    assert shape.shape_type == ShapeType.TEXT_BOX
    assert shape.x == 1.0
    assert shape.text_style
    assert shape.source_anchor.page == 1


def test_parse_pptx_rejects_out_of_range_target_slide(tmp_path: Path) -> None:
    path = tmp_path / "source.pptx"
    _create_pptx(path)

    with pytest.raises(ParserError) as exc_info:
        parse_pptx(
            path,
            document_id="doc_01",
            filename=path.name,
            extracted_text_path="jobs/job_1/parsed/doc_01.json",
            target_slide=3,
        )

    assert exc_info.value.code == "TARGET_SLIDE_OUT_OF_RANGE"


def test_parse_pptx_rejects_corrupted_file(tmp_path: Path) -> None:
    path = tmp_path / "source.pptx"
    path.write_bytes(b"PK\x03\x04 not an office archive")

    with pytest.raises(ParserError) as exc_info:
        parse_pptx(
            path,
            document_id="doc_01",
            filename=path.name,
            extracted_text_path="jobs/job_1/parsed/doc_01.json",
        )

    assert exc_info.value.code == "CORRUPTED_FILE"
