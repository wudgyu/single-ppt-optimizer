"""DOCX parser service."""

import hashlib
from pathlib import Path

from docx import Document
from docx.table import Table

from app.schemas.source import (
    ParsedBlock,
    ParsedTable,
    ParsedTableCell,
    SourceAnchor,
    SourceDocument,
    SourceFileType,
)
from app.services.file_validation import validate_document_file
from app.services.parser_errors import ParserError


def _heading_level(style_name: str) -> int | None:
    normalized = style_name.lower()
    if normalized.startswith("heading "):
        suffix = normalized.rsplit(" ", 1)[-1]
        if suffix.isdigit():
            return int(suffix)
    if style_name.startswith("标题 "):
        suffix = style_name.rsplit(" ", 1)[-1]
        if suffix.isdigit():
            return int(suffix)
    return None


def _parse_table(
    table: Table,
    *,
    document_id: str,
    filename: str,
    table_number: int,
) -> ParsedTable:
    table_id = f"table_{table_number}"
    cells = [
        ParsedTableCell(row=row_index, column=column_index, text=cell.text)
        for row_index, row in enumerate(table.rows)
        for column_index, cell in enumerate(row.cells)
    ]
    first_row_text = " | ".join(cell.text for cell in table.rows[0].cells)
    return ParsedTable(
        table_id=table_id,
        row_count=len(table.rows),
        column_count=len(table.columns),
        cells=cells,
        source_anchor=SourceAnchor(
            document_id=document_id,
            filename=filename,
            table_id=table_id,
            text_quote=first_row_text[:500],
        ),
    )


def parse_docx(
    path: Path,
    *,
    document_id: str,
    filename: str,
    extracted_text_path: str,
    max_size_bytes: int = 100 * 1024 * 1024,
) -> SourceDocument:
    validate_document_file(path, SourceFileType.DOCX, max_size_bytes=max_size_bytes)
    try:
        document = Document(str(path))
    except Exception as exc:
        raise ParserError(
            "CORRUPTED_FILE",
            "DOCX file could not be opened.",
            details={"filename": filename},
        ) from exc

    blocks: list[ParsedBlock] = []
    for paragraph_index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        heading_level = _heading_level(paragraph.style.name if paragraph.style else "")
        blocks.append(
            ParsedBlock(
                block_id=f"docx_paragraph_{paragraph_index}",
                block_type="heading" if heading_level else "paragraph",
                text=text,
                heading_level=heading_level,
                paragraph_index=paragraph_index,
                source_anchor=SourceAnchor(
                    document_id=document_id,
                    filename=filename,
                    paragraph_index=paragraph_index,
                    text_quote=text[:500],
                ),
            )
        )

    for table_number, table in enumerate(document.tables, start=1):
        parsed_table = _parse_table(
            table,
            document_id=document_id,
            filename=filename,
            table_number=table_number,
        )
        blocks.append(
            ParsedBlock(
                block_id=f"docx_{parsed_table.table_id}",
                block_type="table",
                text="\n".join(cell.text for cell in parsed_table.cells),
                table=parsed_table,
                source_anchor=parsed_table.source_anchor,
                limitations=[],
            )
        )

    return SourceDocument(
        document_id=document_id,
        filename=filename,
        file_type=SourceFileType.DOCX,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        page_count=1,
        extracted_text_path=extracted_text_path,
        parsed_blocks=blocks,
    )
