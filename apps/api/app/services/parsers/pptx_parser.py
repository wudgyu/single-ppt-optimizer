"""PPTX parser service."""

import hashlib
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.util import Emu

from app.schemas.source import (
    ParsedChart,
    ParsedChartSeries,
    ParsedShape,
    ParsedSlide,
    ParsedTable,
    ParsedTableCell,
    ShapeType,
    SourceAnchor,
    SourceDocument,
    SourceFileType,
    TextStyle,
)
from app.services.file_validation import validate_document_file
from app.services.parser_errors import ParserError

EMU_PER_INCH = 914400


def _inches(value: int | None) -> float | None:
    return round(Emu(value).inches, 6) if value is not None else None


def _first_text_style(shape: Any) -> TextStyle | None:
    if not getattr(shape, "has_text_frame", False):
        return None
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            font = run.font
            color = None
            try:
                if font.color and font.color.rgb:
                    color = str(font.color.rgb)
            except (AttributeError, TypeError):
                color = None
            return TextStyle(
                font_size_pt=float(font.size.pt) if font.size else None,
                font_name=font.name,
                bold=font.bold,
                color=color,
                alignment=str(paragraph.alignment) if paragraph.alignment else None,
            )
    return None


def _shape_type(shape: Any) -> ShapeType:
    if getattr(shape, "has_table", False):
        return ShapeType.TABLE
    if getattr(shape, "has_chart", False):
        return ShapeType.CHART
    if getattr(shape, "shape_type", None) is not None:
        name = str(shape.shape_type).lower()
        if "group" in name:
            return ShapeType.GROUP
        if "picture" in name:
            return ShapeType.PICTURE
        if "text_box" in name:
            return ShapeType.TEXT_BOX
        if "auto_shape" in name:
            return ShapeType.AUTO_SHAPE
    return ShapeType.UNKNOWN


def _parse_table(shape: Any, anchor: SourceAnchor, table_id: str) -> ParsedTable:
    table = shape.table
    cells: list[ParsedTableCell] = []
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cells.append(
                ParsedTableCell(
                    row=row_index,
                    column=column_index,
                    text=cell.text,
                )
            )
    return ParsedTable(
        table_id=table_id,
        row_count=len(table.rows),
        column_count=len(table.columns),
        cells=cells,
        source_anchor=anchor,
    )


def _parse_chart(
    shape: Any, anchor: SourceAnchor, chart_id: str, limitations: list[str]
) -> ParsedChart | None:
    try:
        chart = shape.chart
        plots = list(chart.plots)
        if not plots:
            return None
        plot = plots[0]
        categories = [str(value) for value in plot.categories]
        series = [
            ParsedChartSeries(
                name=str(series.name),
                values=[float(value) if value is not None else None for value in series.values],
            )
            for series in plot.series
        ]
        return ParsedChart(
            chart_id=chart_id,
            categories=categories,
            series=series,
            accessible=True,
            source_anchor=anchor,
        )
    except Exception:
        limitations.append("Chart data was not accessible; no values were inferred.")
        return None


def _parse_shape(
    shape: Any,
    document_id: str,
    filename: str,
    slide_number: int,
    path: str,
) -> ParsedShape:
    shape_id = f"shape_{shape.shape_id}"
    text = shape.text_frame.text if getattr(shape, "has_text_frame", False) else ""
    anchor = SourceAnchor(
        document_id=document_id,
        filename=filename,
        page=slide_number,
        shape_id=shape_id,
        text_quote=text[:500],
    )
    limitations: list[str] = []
    shape_type = _shape_type(shape)
    tables: list[ParsedTable] = []
    charts: list[ParsedChart] = []

    if shape_type == ShapeType.TABLE:
        tables.append(_parse_table(shape, anchor, f"{shape_id}_table_1"))
    if shape_type == ShapeType.CHART:
        chart = _parse_chart(shape, anchor, f"{shape_id}_chart_1", limitations)
        if chart:
            charts.append(chart)
    if shape_type == ShapeType.UNKNOWN:
        limitations.append("Object type could not be classified; no behavior was inferred.")

    return ParsedShape(
        shape_id=shape_id,
        shape_type=shape_type,
        text=text,
        x=_inches(shape.left),
        y=_inches(shape.top),
        w=_inches(shape.width),
        h=_inches(shape.height),
        rotation_degrees=float(shape.rotation),
        text_style=_first_text_style(shape),
        tables=tables,
        charts=charts,
        source_anchor=anchor,
        limitations=limitations,
    )


def parse_pptx(
    path: Path,
    *,
    document_id: str,
    filename: str,
    extracted_text_path: str,
    max_size_bytes: int = 100 * 1024 * 1024,
    max_slides: int = 100,
    target_slide: int | None = None,
) -> SourceDocument:
    validate_document_file(path, SourceFileType.PPTX, max_size_bytes=max_size_bytes)
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ParserError(
            "CORRUPTED_FILE",
            "PPTX file could not be opened.",
            details={"filename": filename},
        ) from exc

    page_count = len(presentation.slides)
    if page_count == 0:
        raise ParserError(
            "CORRUPTED_FILE", "PPTX file has no slides.", details={"filename": filename}
        )
    if page_count > max_slides:
        raise ParserError(
            "FILE_TOO_LARGE",
            "PPTX slide count exceeds the configured limit.",
            details={"page_count": page_count, "max_slides": max_slides},
        )
    if target_slide is not None and not 1 <= target_slide <= page_count:
        raise ParserError(
            "TARGET_SLIDE_OUT_OF_RANGE",
            "Target slide is outside the source document.",
            details={"target_slide": target_slide, "page_count": page_count},
        )

    slides: list[ParsedSlide] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        shapes = [
            _parse_shape(shape, document_id, filename, slide_number, "shapes")
            for shape in slide.shapes
        ]
        try:
            background = str(slide.background.fill.type)
        except (AttributeError, TypeError):
            background = "unknown"
        slides.append(
            ParsedSlide(
                slide_number=slide_number,
                canvas_width_inches=_inches(presentation.slide_width) or 0,
                canvas_height_inches=_inches(presentation.slide_height) or 0,
                background=background,
                shapes=shapes,
            )
        )

    return SourceDocument(
        document_id=document_id,
        filename=filename,
        file_type=SourceFileType.PPTX,
        sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        page_count=page_count,
        extracted_text_path=extracted_text_path,
        parsed_slides=slides,
    )
